"""
Trade Journal - the single place you record what you actually BUY and SELL
---------------------------------------------------------------------------
Added 2026-10-09. Lives in Trade_Tracker.xlsx, sheet "My Trades".

One row per fill, typed by hand (yellow cells):
    Date | Ticker | Action (BUY / SELL / OPENING) | Call/Put | Strike | Expiry |
    Shares / Contracts | Price (per share) | Fees | Notes

STOCK row  : leave Call/Put, Strike, Expiry blank. Qty = shares.
OPTION row : fill Call/Put (C or P), Strike and Expiry. Qty = contracts,
             Price = premium PER SHARE as your broker quotes it (21.80 =
             $2,180 per contract). Added 2026-10-09 so LEAPs are keyed once
             too; open option positions feed Portfolio_Tracker "My Options".

Everything else is worked out from those rows:
  - current holdings (shares + average cost)   -> Portfolio_Tracker.xlsx "My Portfolio"
  - realized P/L on every SELL                  -> Trade_Tracker.xlsx + the dashboard
  - what you bought each month                  -> the Recommendation Log's "Actual" columns

Cost method: AVERAGE COST (what moomoo and most brokers display). A BUY
raises the share count and re-averages the cost (fees included in cost).
A SELL removes shares at the current average cost; realized P/L =
(sell price - average cost) x shares - fees. Average cost does not change
on a partial sell, and resets once a position is fully closed.

OPENING = a position you already held before the journal started. It works
exactly like a BUY for holdings, but is not counted as "bought this month"
in the Recommendation Log.

DIVIDEND (added 2026-10-21) = cash dividend received. Ticker, Date (pay
date), Shares = shares the dividend was paid on, Price = dividend PER SHARE,
Fees = withholding tax taken by the broker (30% for Malaysian residents on
US dividends) plus any charges. Amount = shares x DPS - fees. It does not
change your holdings or average cost; it is counted as income on the
dashboard's Actual P&L section (month, quarter, year to date).

This module only READS the sheet and computes. Writing the calculated
columns back is done by generate_trade_tracker.py.
"""

import datetime
import os

TRACKER_FILE = "Trade_Tracker.xlsx"
TRADES_SHEET = "My Trades"
OPT_MULTIPLIER = 100  # 1 US equity option contract = 100 shares

# (key, header text) - the header text is what appears in row 1 of the sheet
INPUT_FIELDS = [("date", "Date"), ("ticker", "Ticker"), ("action", "Action"),
                ("cp", "Call/Put"), ("strike", "Strike"), ("expiry", "Expiry"),
                ("qty", "Shares / Contracts"), ("price", "Price (per share)"),
                ("fees", "Fees"), ("notes", "Notes")]
INPUT_HEADERS = [h for _, h in INPUT_FIELDS]
FIELD_BY_HEADER = {h: k for k, h in INPUT_FIELDS}
FIELD_BY_HEADER.update({"Shares": "qty", "Price": "price"})  # first (stock-only) layout
CALC_HEADERS = ["Position", "Company", "Trade Value", "Avg Cost Used", "Realized P/L ($)",
                "Realized P/L (%)", "Held After", "Avg Cost After", "Check"]
ALL_HEADERS = INPUT_HEADERS + CALC_HEADERS
ACTIONS = ("BUY", "SELL", "OPENING", "DIVIDEND")

EPS = 1e-6


def _f(v):
    try:
        if v in (None, ""):
            return None
        x = float(v)
        return None if x != x else x
    except (TypeError, ValueError):
        return None


def parse_date(v):
    """Real Excel date, or text. Text is read day-first (Malaysian format):
    9/10/2026 = 9 Oct 2026. ISO 2026-10-09 also works."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y",
                "%d %B %Y", "%b %d %Y", "%b %d, %Y", "%Y/%m/%d", "%d/%m/%y", "%d-%b-%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _norm_action(v):
    s = str(v or "").strip().upper()
    if s in ("B", "BUY", "BOUGHT", "ADD", "BTO"):
        return "BUY"
    if s in ("S", "SELL", "SOLD", "TRIM", "CLOSE", "STC"):
        return "SELL"
    if s in ("OPENING", "OPEN", "OPENING BALANCE", "OPEN BAL"):
        return "OPENING"
    if s in ("DIVIDEND", "DIVIDENDS", "DIV", "DIV.", "DVD"):
        return "DIVIDEND"
    return None


def norm_cp(v):
    s = str(v or "").strip().upper()
    if s in ("C", "CALL", "CALLS"):
        return "C"
    if s in ("P", "PUT", "PUTS"):
        return "P"
    return None


def _is_ticker(v):
    return bool(v) and isinstance(v, str) and 0 < len(v.strip()) <= 10 and " " not in v.strip()


def position_label(ticker, cp=None, strike=None, expiry=None):
    """NVDA  /  IREN C25 21-Jan-2028"""
    if not cp:
        return ticker
    k = f"{strike:g}" if isinstance(strike, (int, float)) else str(strike)
    e = expiry.strftime("%d-%b-%Y") if isinstance(expiry, datetime.date) else str(expiry or "?")
    return f"{ticker} {cp}{k} {e}"


def header_map(ws):
    """field key -> column number, accepting old header names too."""
    out = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h in FIELD_BY_HEADER and FIELD_BY_HEADER[h] not in out:
            out[FIELD_BY_HEADER[h]] = c
    return out


def read_trades(ws):
    """Rows from the My Trades sheet -> list of dicts (input fields only),
    each tagged with its sheet row number. Blank rows are skipped."""
    col = header_map(ws)
    missing = [h for k, h in INPUT_FIELDS if k not in col and k not in ("cp", "strike", "expiry")]
    if missing:
        raise ValueError(f"'{TRADES_SHEET}' sheet is missing column(s): {missing}")
    out = []
    for r in range(2, ws.max_row + 1):
        get = lambda k: ws.cell(row=r, column=col[k]).value if k in col else None
        ticker, action, qty, price = get("ticker"), get("action"), get("qty"), get("price")
        if not ticker and action in (None, "") and qty in (None, "") and price in (None, ""):
            continue  # empty row
        if isinstance(ticker, str) and not _is_ticker(ticker) and len(ticker) > 10:
            continue  # a note line typed in the Ticker column
        cp_raw, strike_raw, exp_raw = get("cp"), get("strike"), get("expiry")
        is_opt = any(v not in (None, "") for v in (cp_raw, strike_raw, exp_raw))
        out.append({
            "row": r,
            "date_raw": get("date"),
            "date": parse_date(get("date")),
            "ticker": str(ticker).strip().upper() if ticker else None,
            "action_raw": action,
            "action": _norm_action(action),
            "is_option": is_opt,
            "cp_raw": cp_raw,
            "cp": norm_cp(cp_raw) if is_opt else None,
            "strike": _f(strike_raw) if is_opt else None,
            "expiry_raw": exp_raw,
            "expiry": parse_date(exp_raw) if is_opt else None,
            "shares": _f(qty),   # shares for a stock row, contracts for an option row
            "price": _f(price),
            "fees": _f(get("fees")) or 0.0,
            "notes": get("notes"),
        })
    return out


def compute(trades, today=None, dividends_out=None):
    """Walk the trades in date order and work out per-row results, current
    positions and realized P/L. Stocks and each option contract (ticker +
    call/put + strike + expiry) are separate positions.

    Returns (results_by_row, positions, sells)
      results_by_row[row] = {label, trade_value, avg_cost_used, realized, realized_pct,
                             shares_after, avg_after, check}
      positions[key]      = {kind 'STOCK'/'OPTION', ticker, cp, strike, expiry, label,
                             shares (qty: shares or contracts), avg_cost (per share),
                             cost_basis ($), realized, first_date, last_date,
                             n_buys, n_sells, notes}   (closed ones have shares == 0)
      sells               = list of dicts, one per valid SELL
    dividends_out, if a list is passed, receives one dict per valid DIVIDEND row
    (DIVIDEND rows never change holdings, cost or realized P/L on sells).
    """
    today = today or datetime.date.today()
    results, positions, sells = {}, {}, []
    last_row_of = {}

    def sort_key(t):
        # OPENING balances always come first (they describe what you held
        # before the journal started, whatever date is typed on them);
        # undated BUY/SELL rows go last, in sheet order.
        if t["action"] == "OPENING":
            return (0, datetime.date.min, t["row"])
        return (1, t["date"] or datetime.date.max, t["row"])

    for t in sorted(trades, key=sort_key):
        problems = []
        if not _is_ticker(t["ticker"] or ""):
            problems.append("Ticker missing")
        if not t["action"]:
            problems.append("Action must be BUY, SELL, OPENING or DIVIDEND")
        unit = "Contracts" if t["is_option"] else "Shares"
        if not t["shares"] or t["shares"] <= 0:
            problems.append(f"{unit} must be > 0")
        if t["price"] is None or t["price"] < 0:
            problems.append("Price missing")
        if t["is_option"]:
            if not t["cp"]:
                problems.append("Call/Put must be C or P")
            if not t["strike"]:
                problems.append("Strike missing")
            if not t["expiry"]:
                problems.append("Expiry not a date")
        if t["date"] is None and t["action"] != "OPENING":
            problems.append("Date missing/unreadable")  # warning only, row still counts
        hard = [p for p in problems if not p.startswith("Date")]
        if hard:
            results[t["row"]] = {"check": "; ".join(problems)}
            continue

        tk = t["ticker"]
        if t["action"] == "DIVIDEND":
            if t["is_option"]:
                results[t["row"]] = {"check": "DIVIDEND rows are for stocks - leave Call/Put, Strike, Expiry blank"}
                continue
            sp = positions.get(("STOCK", tk))
            gross = t["shares"] * t["price"]
            amount = gross - t["fees"]
            held = sp["shares"] if sp else 0.0
            cost_ps = sp["cost_basis"] / sp["shares"] if sp and sp["shares"] > EPS else None
            if held + EPS < t["shares"]:
                problems.append(f"Dividend on {t['shares']:g} shares but {held:g} held on that date - "
                                "check the date or shares (still counted)")
            res = {"label": f"{tk} dividend", "trade_value": round(gross, 2), "realized": round(amount, 2),
                   "realized_pct": round(amount / (cost_ps * t["shares"]), 4) if cost_ps else None,
                   "shares_after": round(held, 6),
                   "avg_after": round(cost_ps, 4) if cost_ps else None,
                   "check": "; ".join(problems) if problems else "OK"}
            results[t["row"]] = res
            if dividends_out is not None:
                dividends_out.append({"row": t["row"], "date": t["date"], "ticker": tk, "shares": t["shares"],
                                      "dps": t["price"], "fees": t["fees"], "gross": round(gross, 2),
                                      "amount": round(amount, 2), "avg_cost": round(cost_ps, 4) if cost_ps else None,
                                      "yield_on_cost": res["realized_pct"], "notes": t["notes"]})
            continue
        if t["is_option"]:
            key = ("OPTION", tk, t["cp"], t["strike"], t["expiry"])
            mult = OPT_MULTIPLIER
        else:
            key = ("STOCK", tk)
            mult = 1
        label = position_label(tk, t["cp"], t["strike"], t["expiry"])
        p = positions.setdefault(key, {"kind": key[0], "ticker": tk, "cp": t["cp"],
                                       "strike": t["strike"], "expiry": t["expiry"], "label": label,
                                       "mult": mult, "shares": 0.0, "cost_basis": 0.0, "realized": 0.0,
                                       "first_date": None, "last_date": None,
                                       "n_buys": 0, "n_sells": 0, "notes": None})
        value = t["shares"] * t["price"] * mult
        res = {"label": label, "trade_value": round(value, 2)}

        if t["action"] in ("BUY", "OPENING"):
            if p["shares"] <= EPS:   # position goes from zero to held: remember when (for "Opened" on the dashboard)
                p["opened"] = "carried over" if t["action"] == "OPENING" else t["date"]
            p["shares"] += t["shares"]
            p["cost_basis"] += value + t["fees"]
            p["n_buys"] += 1
            if t["date"] and (p["first_date"] is None or t["date"] < p["first_date"]):
                p["first_date"] = t["date"]
            if p["notes"] is None and isinstance(t["notes"], str) and t["notes"].strip():
                p["notes"] = t["notes"].strip()
        else:  # SELL
            if t["shares"] > p["shares"] + EPS:
                problems.append(f"Selling {t['shares']:g} {unit.lower()} but only {p['shares']:g} held on that "
                                "date - check the date (and strike/expiry for an option) or add the missing BUY")
                results[t["row"]] = {"label": label, "check": "; ".join(problems), "trade_value": round(value, 2)}
                continue
            avg_total = p["cost_basis"] / p["shares"] if p["shares"] > EPS else 0.0  # $ per share/contract
            cost_out = avg_total * t["shares"]
            realized = value - t["fees"] - cost_out
            p["shares"] -= t["shares"]
            p["cost_basis"] -= cost_out
            if p["shares"] <= EPS:  # fully closed - reset cleanly (no float dust)
                p["shares"], p["cost_basis"] = 0.0, 0.0
            p["realized"] += realized
            p["n_sells"] += 1
            avg_ps = avg_total / mult  # per share, as brokers quote it
            res.update({"avg_cost_used": round(avg_ps, 4), "realized": round(realized, 2),
                        "realized_pct": round(realized / cost_out, 4) if cost_out else None})
            sells.append({"row": t["row"], "date": t["date"], "ticker": tk, "kind": key[0],
                          "label": label, "cp": t["cp"], "strike": t["strike"], "expiry": t["expiry"],
                          "shares": t["shares"], "price": t["price"], "fees": t["fees"],
                          "avg_cost": round(avg_ps, 4), "proceeds": round(value - t["fees"], 2),
                          "cost": round(cost_out, 2), "realized": round(realized, 2),
                          "realized_pct": round(realized / cost_out, 4) if cost_out else None,
                          "closed_position": p["shares"] == 0.0,
                          "opened": p.get("opened"),
                          "notes": t["notes"]})
        if t["date"]:
            p["last_date"] = t["date"] if p["last_date"] is None else max(p["last_date"], t["date"])
        res["shares_after"] = round(p["shares"], 6)
        res["avg_after"] = round(p["cost_basis"] / p["shares"] / mult, 4) if p["shares"] > EPS else None
        res["check"] = "; ".join(problems) if problems else "OK"
        results[t["row"]] = res
        last_row_of[key] = t["row"]

    for key, p in positions.items():
        p["avg_cost"] = round(p["cost_basis"] / p["shares"] / p["mult"], 4) if p["shares"] > EPS else 0.0
        p["shares"] = round(p["shares"], 6)
        p["cost_basis"] = round(p["cost_basis"], 2)
        p["realized"] = round(p["realized"], 2)
        # An option still open after its expiry date needs closing out
        if p["kind"] == "OPTION" and p["shares"] > 0 and p["expiry"] and p["expiry"] < today:
            r = last_row_of.get(key)
            if r in results:
                msg = ("EXPIRED " + p["expiry"].strftime("%d-%b-%Y") + " but still open - add a SELL "
                       "(price 0 if it expired worthless)")
                c = results[r].get("check", "OK")
                results[r]["check"] = msg if c == "OK" else c + "; " + msg
    return results, positions, sells


def stock_positions(positions):
    """ticker -> position, stocks only."""
    return {p["ticker"]: p for p in positions.values() if p["kind"] == "STOCK"}


def option_positions(positions, open_only=True):
    out = [p for p in positions.values() if p["kind"] == "OPTION" and (p["shares"] > 0 or not open_only)]
    return sorted(out, key=lambda p: (p["ticker"], p["expiry"] or datetime.date.max, p["strike"] or 0))


def monthly_buys(trades):
    """(YYYY-MM, ticker) -> {shares, avg_price, first_date} from stock BUY rows
    only (OPENING and options excluded) - feeds the Recommendation Log's
    Actual columns."""
    agg = {}
    for t in trades:
        if (t["action"] != "BUY" or t["is_option"] or not t["date"] or not t["shares"]
                or t["price"] is None or not t["ticker"]):
            continue
        k = (t["date"].strftime("%Y-%m"), t["ticker"])
        a = agg.setdefault(k, {"shares": 0.0, "cost": 0.0, "first_date": t["date"]})
        a["shares"] += t["shares"]
        a["cost"] += t["shares"] * t["price"]
        a["first_date"] = min(a["first_date"], t["date"])
    for a in agg.values():
        a["avg_price"] = round(a["cost"] / a["shares"], 4) if a["shares"] else None
    return agg


def load(path=TRACKER_FILE):
    """Convenience for other scripts. Returns (has_journal, trades, positions, sells).
    has_journal is False when the file/sheet doesn't exist or has no trade
    rows yet - callers then fall back to their old behaviour."""
    if not os.path.exists(path):
        return False, [], {}, []
    from openpyxl import load_workbook
    wb = load_workbook(path)  # inputs are typed values, not formulas - no data_only needed
    if TRADES_SHEET not in wb.sheetnames:
        return False, [], {}, []
    trades = read_trades(wb[TRADES_SHEET])
    if not trades:
        return False, [], {}, []
    _, positions, sells = compute(trades)
    return True, trades, positions, sells


def load_full(path=TRACKER_FILE):
    """Like load(), plus dividends: (has_journal, trades, positions, sells, dividends)."""
    if not os.path.exists(path):
        return False, [], {}, [], []
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if TRADES_SHEET not in wb.sheetnames:
        return False, [], {}, [], []
    trades = read_trades(wb[TRADES_SHEET])
    if not trades:
        return False, [], {}, [], []
    divs = []
    _, positions, sells = compute(trades, dividends_out=divs)
    return True, trades, positions, sells, divs
